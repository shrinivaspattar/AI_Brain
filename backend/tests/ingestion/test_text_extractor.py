from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest

from app.ingestion.text_extractor import extract_text


def test_extract_text_reads_plain_text_files(tmp_path: Path) -> None:
    file_path = tmp_path / "notes.txt"
    file_path.write_text("hello AI_Brain")

    assert extract_text(file_path) == "hello AI_Brain"


def test_extract_text_reads_unlisted_extensions_as_plain_text(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "data.csv"
    file_path.write_text("a,b,c\n1,2,3")

    assert extract_text(file_path) == "a,b,c\n1,2,3"


def test_extract_text_raises_for_missing_file(tmp_path: Path) -> None:
    with pytest.raises(OSError):
        extract_text(tmp_path / "missing.txt")


def test_extract_text_raises_for_binary_content(tmp_path: Path) -> None:
    file_path = tmp_path / "image.bin"
    file_path.write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\xff\xfe")

    with pytest.raises(UnicodeDecodeError):
        extract_text(file_path)


def test_extract_text_joins_pdf_pages() -> None:
    page_one = MagicMock()
    page_one.extract_text.return_value = "page one"

    page_two = MagicMock()
    page_two.extract_text.return_value = "page two"

    with patch("app.ingestion.text_extractor.pypdf.PdfReader") as reader_class:
        reader_class.return_value.pages = [page_one, page_two]

        result = extract_text(Path("document.pdf"))

    assert result == "page one\npage two"
    reader_class.assert_called_once_with("document.pdf")


def test_extract_text_handles_pdf_pages_with_no_extractable_text() -> None:
    page = MagicMock()
    page.extract_text.return_value = None

    with patch("app.ingestion.text_extractor.pypdf.PdfReader") as reader_class:
        reader_class.return_value.pages = [page]

        result = extract_text(Path("scanned.pdf"))

    assert result == ""


def test_extract_text_reads_docx_paragraphs(tmp_path: Path) -> None:
    file_path = tmp_path / "notes.docx"

    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(str(file_path))

    assert extract_text(file_path) == "First paragraph.\nSecond paragraph."


def test_extract_text_handles_pdf_renamed_with_conflict_suffix() -> None:
    # Real-world pattern: a bulk conflict-resolution/dedup rename appends
    # "_<timestamp>" directly after the extension, no separating dot -
    # e.g. "Sept-Oct-2022 - 2nd Sem M Tech VLSI-min.pdf_1768918262".
    page = MagicMock()
    page.extract_text.return_value = "exam questions"

    with patch("app.ingestion.text_extractor.pypdf.PdfReader") as reader_class:
        reader_class.return_value.pages = [page]

        result = extract_text(Path("question-paper.pdf_1768918262"))

    assert result == "exam questions"
    reader_class.assert_called_once_with("question-paper.pdf_1768918262")


def test_extract_text_handles_docx_renamed_with_conflict_suffix(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "notes.docx_1768918262"

    document = docx.Document()
    document.add_paragraph("Renamed but still a real docx.")
    document.save(str(file_path))

    assert extract_text(file_path) == "Renamed but still a real docx."


def test_extract_text_does_not_misfire_on_unrelated_extension_prefix(
    tmp_path: Path,
) -> None:
    # Guard against overly broad matching: ".pdfx" or similar must not be
    # treated as a renamed PDF.
    file_path = tmp_path / "notes.pdfx"
    file_path.write_text("just text, not a pdf")

    assert extract_text(file_path) == "just text, not a pdf"


def test_extract_text_reads_pptx_slide_text(tmp_path: Path) -> None:
    import pptx

    file_path = tmp_path / "deck.pptx"

    presentation = pptx.Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, pptx.util.Inches(4), pptx.util.Inches(2))
    box.text_frame.text = "First slide title"

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    box2 = slide2.shapes.add_textbox(0, 0, pptx.util.Inches(4), pptx.util.Inches(2))
    box2.text_frame.text = "Second slide content"

    presentation.save(str(file_path))

    result = extract_text(file_path)

    assert "First slide title" in result
    assert "Second slide content" in result


def test_extract_text_handles_pptx_renamed_with_conflict_suffix(
    tmp_path: Path,
) -> None:
    import pptx

    file_path = tmp_path / "deck.pptx_1768918262"

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, pptx.util.Inches(4), pptx.util.Inches(2))
    box.text_frame.text = "Renamed but still a real pptx"
    presentation.save(str(file_path))

    assert "Renamed but still a real pptx" in extract_text(file_path)


def test_extract_text_reads_xlsx_cell_values(tmp_path: Path) -> None:
    import openpyxl

    file_path = tmp_path / "sheet.xlsx"

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["name", "score"])
    sheet.append(["Alice", 95])
    workbook.save(str(file_path))

    result = extract_text(file_path)

    assert "name\tscore" in result
    assert "Alice\t95" in result


def test_extract_text_handles_xlsx_renamed_with_conflict_suffix(
    tmp_path: Path,
) -> None:
    import openpyxl

    file_path = tmp_path / "sheet.xlsx_1768918262"

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["renamed", "still works"])
    workbook.save(str(file_path))

    assert "renamed\tstill works" in extract_text(file_path)


def test_extract_text_removes_nul_characters(tmp_path: Path) -> None:
    path = tmp_path / "with_nul.txt"
    path.write_text("before\x00after\x00", encoding="utf-8")

    assert extract_text(path) == "beforeafter"


def test_extract_text_reads_utf16_with_byte_order_mark(tmp_path: Path) -> None:
    path = tmp_path / "windows.txt"
    path.write_bytes("héllo wörld".encode("utf-16"))

    assert extract_text(path) == "héllo wörld"


def test_extract_text_still_rejects_non_utf8_without_a_byte_order_mark(tmp_path: Path) -> None:
    path = tmp_path / "cp1252.txt"
    path.write_bytes(b"20\xb0C")

    with pytest.raises(UnicodeDecodeError):
        extract_text(path)


# --- Obsidian-aware .md handling ------------------------------------------

def test_extract_text_strips_frontmatter_and_shows_title_and_tags(tmp_path: Path) -> None:
    path = tmp_path / "note.md"
    path.write_text(
        '---\ntitle: "My Note"\ntags:\n  - work\n  - ideas\n---\nThe actual body text.'
    )

    result = extract_text(path)

    assert result == "Title: My Note\nTags: work, ideas\n\nThe actual body text."


def test_extract_text_surfaces_web_clipper_style_frontmatter_fields(tmp_path: Path) -> None:
    path = tmp_path / "clipped.md"
    path.write_text(
        "---\n"
        'title: "Some Article"\n'
        'source: "https://example.com/article"\n'
        "author: Jane Doe\n"
        "published: 2026-05-01\n"
        "---\n"
        "Article body here."
    )

    result = extract_text(path)

    assert result == (
        "Title: Some Article\n"
        "Source: https://example.com/article\n"
        "Author: Jane Doe\n"
        "Published: 2026-05-01\n\n"
        "Article body here."
    )


def test_extract_text_preserves_arbitrary_frontmatter_fields(tmp_path: Path) -> None:
    """Regression: an earlier version only surfaced a curated whitelist of
    fields (title/source/author/published/created/description/tags),
    silently dropping anything else - found against a real ingested note
    with `date`/`status` fields. Every present field must survive."""
    path = tmp_path / "note.md"
    path.write_text(
        "---\n"
        "title: Speed Math\n"
        "date: 2026-06-28\n"
        "status: active\n"
        "---\n"
        "Body text."
    )

    result = extract_text(path)

    assert result == "Title: Speed Math\nDate: 2026-06-28\nStatus: active\n\nBody text."


def test_extract_text_leaves_body_unchanged_when_no_frontmatter(tmp_path: Path) -> None:
    path = tmp_path / "plain.md"
    path.write_text("Just a plain note, no header.")

    assert extract_text(path) == "Just a plain note, no header."


def test_extract_text_tolerates_malformed_frontmatter(tmp_path: Path) -> None:
    path = tmp_path / "broken.md"
    path.write_text("---\nnot: [valid: yaml: at all\n---\nBody text.")

    # Malformed frontmatter must never break ingestion - the whole text is
    # kept as-is (including the unparsed header) rather than raising.
    result = extract_text(path)

    assert "Body text." in result


def test_extract_text_rewrites_wikilinks_to_plain_text(tmp_path: Path) -> None:
    path = tmp_path / "linked.md"
    path.write_text("See [[Project Ideas]] and [[Project Ideas|my ideas]] for more.")

    assert extract_text(path) == "See Project Ideas and my ideas for more."


def test_extract_text_drops_embeds(tmp_path: Path) -> None:
    path = tmp_path / "embed.md"
    path.write_text("Some text ![[diagram.png]] more text.")

    assert extract_text(path) == "Some text  more text."


def test_extract_text_handles_md_conflict_suffix(tmp_path: Path) -> None:
    path = tmp_path / "note.md_1768918262"
    path.write_text("---\ntitle: X\n---\nBody.")

    assert extract_text(path) == "Title: X\n\nBody."
